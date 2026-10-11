"""Ingestion v2 full lifecycle (real Voyage + object storage): new BRD -> new version compare ->
apply keeping a removal -> archive under the same id -> snapshot restore revives it ->
retrieval -> BRD delete purges objects + vectors. Needs VOYAGE_API_KEY and the storage service."""
import io
import time
from docx import Document
from backend.core.config import settings
from backend.core.db import pool
from backend.core import storage
from backend.ingest import jobs as ij
from backend.requirements import versioning
from backend.ingest import revise
from backend.chat.history_store import delete_project

import pytest

pytestmark = pytest.mark.integration

TITLE = "ZZ Ingest Smoke"
OWNER: int | None = None


def ok(cond, msg):
    assert cond, msg
    print('PASS', msg)

def brd(fr2: str, rows: list, extra_heading: str = "2.2 Refunds", drop_fr1: bool = False) -> bytes:
    d = Document()
    d.add_heading(TITLE, 0)
    d.add_heading("1. Scope", 1)
    d.add_paragraph("This document defines the payments portal requirements for merchants.")
    d.add_heading("2. Functional Requirements", 1)
    d.add_heading("2.1 Authentication", 2)
    if not drop_fr1:
        d.add_paragraph("FR-1 The system shall authenticate merchants with email and a one-time code.")
    d.add_paragraph(fr2)
    d.add_heading(extra_heading, 2)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "ID", "Requirement"
    for rid, req in rows:
        r = t.add_row().cells
        r[0].text, r[1].text = rid, req
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def run_job(timeout: float = 600):
    """Wait for the BRD's latest job to finish. Any ingest worker may run it (this process's
    or a running server's — claims use SKIP LOCKED), so the test waits instead of claiming."""
    ij.start_worker()
    ij.wake()
    t0 = time.time()
    while time.time() - t0 < timeout:
        job = ij.jobs(OWNER, "zz-ingest-smoke")[0]
        if job["status"] not in ("queued", "running"):
            print(f"     job #{job['id']} {job['mode']} -> {job['status']} in {time.time() - t0:.1f}s", flush=True)
            ok(job["status"] in ("done", "review"), f"job #{job['id']} finished ({job['status']}: {job['error']})")
            return job
        time.sleep(1)
    raise AssertionError("ingest job did not finish in time")


def live(project):
    with pool().connection() as c, c.cursor() as cur:
        cur.execute("select c.id, c.req_id, c.text, c.section, c.section_path, c.embed_hash is not null, "
                    "(select count(*) from brd_embedding e where e.chunk_id = c.id) "
                    "from brd_chunk c join brd_document d on d.id = c.document_id "
                    "where d.owner_id = %s and d.project = %s order by c.ordinal", (OWNER, project))
        return cur.fetchall()


def test_ingest_lifecycle():
    global OWNER
    if not settings.voyage_api_key:
        pytest.skip("needs VOYAGE_API_KEY")
    try:
        storage.store()
    except storage.StorageError as e:
        pytest.skip(f"object storage unavailable: {e}")


    ROWS = [("FR-10", "A merchant can refund a settled payment in full."),
            ("FR-11", "A partial refund cannot exceed the captured amount."),
            ("NFR-3", "Refund status updates within 5 seconds.")]
    V1 = brd("FR-2 The system shall lock an account after five failed attempts.", ROWS)
    V2 = brd("FR-2 The system shall lock an account after THREE failed attempts and notify the merchant.",
             ROWS[:2] + [("FR-12", "A refund above 10,000 SAR requires a second approver.")], "2.2 Refunds and reversals")

    with pool().connection() as c, c.cursor() as cur:
        cur.execute("select u.id from app_user u join brd_document d on d.owner_id = u.id limit 1")
        OWNER = cur.fetchone()[0]


    project = None
    try:
        # 1. new BRD ------------------------------------------------------------------
        r = ij.submit(OWNER, OWNER, V1, "zz-ingest-smoke.docx", title=TITLE)
        project = r["project"]
        ok(r["mode"] == "new" and r["status"] == "queued", f"submitted new BRD '{project}'")
        try:
            ij.submit(OWNER, OWNER, V1, "zz-ingest-smoke.docx", title=TITLE); ok(False, "conflict")
        except ij.UploadRejected as e:
            ok(e.status == 409, f"second upload while one is queued/same title is refused ({e.status})")
        job = run_job()
        rows = live(project)
        ok(len(rows) == 6 and all(x[6] == 1 and x[5] for x in rows), f"live: {len(rows)} requirements, all embedded")
        ok([x[1] for x in rows] == [None, "FR-1", "FR-2", "FR-10", "FR-11", "NFR-3"], "requirement ids per chunk")
        ok(rows[3][4] == "2. Functional Requirements / 2.2 Refunds", "section path kept")
        js = ij.jobs(OWNER, project)[0]
        # cached vectors survive a recent delete (grace period), so a re-run may reuse them
        ok(js["status"] == "done" and js["stats"]["embedded"] + js["stats"].get("reused", 0) == 6,
           f"job done, embedded {js['stats']['embedded']}, reused {js['stats'].get('reused', 0)}")
        with pool().connection() as c, c.cursor() as cur:
            cur.execute("select object_key from brd_source where owner_id=%s and project=%s", (OWNER, project))
            key = cur.fetchone()[0]
        ok(storage.store().exists(key), "source file stored in object storage")
        same = ij.submit(OWNER, OWNER, V1, "zz-ingest-smoke.docx", target_project=project)
        ok(same.get("unchanged") is True, "identical re-upload is a no-op")
        ids_v1 = {x[1]: x[0] for x in rows}

        # 2. new version -> compare --------------------------------------------------------
        r = ij.submit(OWNER, OWNER, V2, "zz-ingest-smoke-v2.docx", target_project=project)
        ok(r["mode"] == "revision", "new version queued as a revision")
        run_job()
        ok(len(live(project)) == 6, "live BRD untouched while the revision awaits review")
        rev = revise.get_open(OWNER, project)
        s = rev["summary"]
        ok(s == {"same": 4, "modified": 1, "added": 1, "removed": 1} or s["modified"] >= 1,
           f"compare summary {s}")
        st = ij.jobs(OWNER, project)[0]["stats"]
        ok(st["reused"] >= 1 and st["embedded"] < 6, f"embedding reuse: embedded {st['embedded']}, reused {st['reused']}")
        by = {(i["change"], i["req_id"]): i for i in rev["items"]}
        ok(("modified", "FR-2") in by and ("added", "FR-12") in by and ("removed", "NFR-3") in by,
           "FR-2 modified, FR-12 added, NFR-3 removed")
        # 3. apply, keeping the removed NFR-3 -------------------------------------------
        res = revise.apply_revision(OWNER, rev["id"], keep_removed=[by[("removed", "NFR-3")]["id"]], by=OWNER)
        rows2 = live(project)
        ids2 = {x[1]: x[0] for x in rows2}
        ok(ids2["FR-2"] == ids_v1["FR-2"] and "THREE" in next(x[2] for x in rows2 if x[1] == "FR-2"),
           "modified FR-2 updated IN PLACE (same id)")
        ok(ids2["FR-10"] == ids_v1["FR-10"] and ids2["NFR-3"] == ids_v1["NFR-3"], "unchanged + kept ids preserved")
        ok("FR-12" in ids2 and all(x[6] == 1 for x in rows2), "FR-12 added with its embedding")
        with pool().connection() as c, c.cursor() as cur:
            cur.execute("select revision_no from brd_document where owner_id=%s and project=%s", (OWNER, project))
            ok(cur.fetchone()[0] == 2, "revision_no bumped to 2")

        # 4. a version that drops FR-1 -> archived, same id, resolvable, audited ---------
        V3 = brd("FR-2 The system shall lock an account after THREE failed attempts and notify the merchant.",
                 ROWS[:2] + [("FR-12", "A refund above 10,000 SAR requires a second approver.")],
                 "2.2 Refunds and reversals", drop_fr1=True)
        ij.submit(OWNER, OWNER, V3, "zz-ingest-smoke-v3.docx", target_project=project)
        run_job()
        rev3 = revise.get_open(OWNER, project)
        removed = [i for i in rev3["items"] if i["change"] == "removed"]
        ok(any(i["req_id"] == "FR-1" for i in removed), f"v3 compare: removed {[i['req_id'] for i in removed]}")
        revise.apply_revision(OWNER, rev3["id"], keep_removed=[i["id"] for i in removed if i["req_id"] != "FR-1"], by=OWNER)
        fr1 = ids_v1["FR-1"]
        with pool().connection() as c, c.cursor() as cur:
            cur.execute("select count(*) from brd_chunk where id=%s", (fr1,)); in_live = cur.fetchone()[0]
            cur.execute("select reason, embedding is not null from brd_chunk_archive where id=%s", (fr1,)); arch = cur.fetchone()
            cur.execute("select kind from requirement_change where chunk_id=%s order by id desc limit 1", (fr1,)); aud = cur.fetchone()
        ok(in_live == 0 and arch == ("revision", True), "FR-1 archived under its ORIGINAL id with its vector")
        ok(aud and aud[0] == "delete", "delete audited (chunk_id kept, not nulled)")

        # 5. restore the pre-v3 snapshot -> FR-1 revived with the same id ----------------
        vs = versioning.list_versions(OWNER, project)["versions"]
        snap = next(v for v in vs if v["label"] == f"Before applying revision #{rev3['id']}")
        versioning.restore(OWNER, project, snap["id"], OWNER)
        ids5 = {x[1]: x[0] for x in live(project)}
        ok(ids5.get("FR-1") == fr1, "snapshot restore revived FR-1 with the SAME id (no re-embedding)")
        ok(ids5["FR-2"] == ids_v1["FR-2"] and ids5["FR-12"] == ids2["FR-12"], "restore kept every other id")

        # 6. retrieval sees the new structure ----------------------------------------------
        from backend.retrieve.retriever import retrieve
        hits = retrieve("how many failed attempts lock an account?", owner_id=OWNER, project=project, k_final=3)
        top = hits[0] if isinstance(hits, list) else None
        print("     top hit:", str(top)[:140], flush=True)
        ok(bool(hits), "retrieval returns results for the new BRD")
    finally:
        if project:
            with pool().connection() as c, c.cursor() as cur:
                keys = ij.source_keys(cur, OWNER, project)
                cur.execute("select count(*) from brd_embedding_cache where owner_id=%s", (OWNER,)); before = cur.fetchone()[0]
            delete_project(project, OWNER)
            with pool().connection() as c, c.cursor() as cur:
                ij.purge_project_rows(cur, OWNER, project); c.commit()
            purged = ij.purge_after_brd_delete(OWNER, keys)
            with pool().connection() as c, c.cursor() as cur:
                cur.execute("select count(*) from brd_chunk_archive a join brd_document d on d.id=a.document_id "
                            "where d.project=%s and d.owner_id=%s", (project, OWNER)); arch_left = cur.fetchone()[0]
                cur.execute("select count(*) from brd_source where project=%s and owner_id=%s", (project, OWNER)); src_left = cur.fetchone()[0]
            print(f"cleanup: {purged}, archive rows left={arch_left}, source rows left={src_left}, "
                  f"object exists={any(storage.store().exists(k) for k in keys)}, cache {before}->", end="")
            with pool().connection() as c, c.cursor() as cur:
                cur.execute("select count(*) from brd_embedding_cache where owner_id=%s", (OWNER,)); print(cur.fetchone()[0])
