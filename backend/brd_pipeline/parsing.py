"""Stage 1 — Parsing. Pure, no DB side effects.

Detect file type by extension and dispatch to the right parser:
  - .pdf  -> PyMuPDF (fitz), text extracted per page in reading order.
  - .docx -> python-docx, paragraphs AND tables walked in document order so
             requirement tables (common in BRDs) keep their place in the flow.

Each parser returns a `ParsedDocument`: raw text plus basic source metadata
(page/paragraph count) and the resolved source filename.
"""
from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Iterator, Union

logger = logging.getLogger(__name__)

# A path (str / os.PathLike) or an already-open binary file-like object.
Source = Union[str, os.PathLike, BinaryIO]

_SUPPORTED = {".pdf", ".docx"}


class UnsupportedFileTypeError(ValueError):
    """Raised when the file extension is not .pdf or .docx."""


class DocumentParseError(RuntimeError):
    """Raised when a supported file cannot be parsed (corrupt/empty/etc.)."""


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    source_filename: str
    file_type: str          # 'pdf' | 'docx'
    page_count: int         # 0 for docx (no fixed pages)
    paragraph_count: int    # 0 for pdf; block count for docx


def _resolve_name(source: Source, filename: str | None) -> str:
    """Determine a filename (for extension detection + metadata)."""
    if filename:
        return os.path.basename(filename)
    if isinstance(source, (str, os.PathLike)):
        return os.path.basename(os.fspath(source))
    raise UnsupportedFileTypeError(
        "A file-like object was given without a `filename`; cannot detect type."
    )


def parse_document(source: Source, *, filename: str | None = None) -> ParsedDocument:
    """Dispatch to the correct parser by extension. Reject unsupported types."""
    name = _resolve_name(source, filename)
    ext = Path(name).suffix.lower()
    if ext not in _SUPPORTED:
        raise UnsupportedFileTypeError(
            f"Unsupported file type {ext!r} for {name!r}. Supported: {sorted(_SUPPORTED)}."
        )
    logger.info("parse_document: dispatching %s (%s)", name, ext)
    if ext == ".pdf":
        return parse_pdf(source, filename=name)
    return parse_docx(source, filename=name)


def parse_pdf(source: Source, *, filename: str) -> ParsedDocument:
    """Extract text from a PDF with PyMuPDF, preserving page + block reading order."""
    import fitz  # PyMuPDF; imported lazily so pure-text tests need no native dep

    try:
        if isinstance(source, (str, os.PathLike)):
            doc = fitz.open(os.fspath(source))
        else:
            doc = fitz.open(stream=source.read(), filetype="pdf")
    except Exception as err:  # noqa: BLE001 - surface a clear, typed error
        raise DocumentParseError(f"Failed to open PDF {filename!r}: {err}") from err

    pages: list[str] = []
    try:
        with doc:
            for page in doc:
                # "blocks" returns layout blocks in reading order: (x0,y0,x1,y1,text,no,type)
                blocks = page.get_text("blocks")
                blocks.sort(key=lambda b: (round(b[1], 1), round(b[0], 1)))  # top-to-bottom, left-to-right
                page_text = "\n\n".join(b[4].strip() for b in blocks if b[4] and b[4].strip())
                if page_text:
                    pages.append(page_text)
            page_count = doc.page_count
    except Exception as err:  # noqa: BLE001
        raise DocumentParseError(f"Failed to read PDF {filename!r}: {err}") from err

    text = "\n\n".join(pages)
    if not text.strip():
        raise DocumentParseError(f"PDF {filename!r} yielded no extractable text.")
    logger.info("parse_pdf: %s -> %d pages", filename, page_count)
    return ParsedDocument(
        text=text,
        source_filename=filename,
        file_type="pdf",
        page_count=page_count,
        paragraph_count=0,
    )


def _iter_docx_blocks(document) -> Iterator[str]:
    """Yield paragraph and table-row text in true document order.

    python-docx exposes paragraphs and tables as separate collections, which
    loses their interleaving. Walking the body's XML children restores order so
    a requirements table isn't hoisted above the paragraph that introduces it.
    """
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document).text
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row in table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def parse_docx(source: Source, *, filename: str) -> ParsedDocument:
    """Extract text from a .docx with python-docx, keeping paragraph/table order."""
    from docx import Document

    try:
        document = Document(source)  # accepts a path or a file-like object
    except Exception as err:  # noqa: BLE001
        raise DocumentParseError(f"Failed to open DOCX {filename!r}: {err}") from err

    blocks = [b.strip() for b in _iter_docx_blocks(document)]
    blocks = [b for b in blocks if b]
    text = "\n\n".join(blocks)
    if not text.strip():
        raise DocumentParseError(f"DOCX {filename!r} yielded no extractable text.")
    logger.info("parse_docx: %s -> %d blocks", filename, len(blocks))
    return ParsedDocument(
        text=text,
        source_filename=filename,
        file_type="docx",
        page_count=0,
        paragraph_count=len(blocks),
    )
